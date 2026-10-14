"""Celery tasks for asynchronous resume parsing and scoring."""
import logging
import os

from celery import shared_task
from celery.signals import worker_process_init
from django.db import OperationalError, transaction
from docx import Document
from pdfminer.high_level import extract_text as pdf_extract_text

from .models import Application
from .utils import calculate_match_score, extract_skills, get_nlp

logger = logging.getLogger(__name__)


class ResumeParsingError(Exception):
    """Permanent, non-retryable problem with the uploaded file."""


@worker_process_init.connect
def _preload_nlp(**_kwargs):
    """Load spaCy once per worker process instead of on the first task."""
    get_nlp()


def _docx_text(fh) -> str:
    doc = Document(fh)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def extract_text_from_resume(field_file) -> str:
    """Extract plain text from a PDF or DOCX FieldFile (works with any storage backend)."""
    ext = os.path.splitext(field_file.name)[1].lower()
    if ext not in (".pdf", ".docx"):
        raise ResumeParsingError(f"Unsupported file type: {ext or 'unknown'}")
    try:
        with field_file.open("rb") as fh:
            text = pdf_extract_text(fh) if ext == ".pdf" else _docx_text(fh)
    except Exception as exc:  # corrupt/encrypted files raise assorted library errors
        raise ResumeParsingError(f"Could not read {ext} file: {exc}") from exc

    text = (text or "").strip()
    if not text:
        raise ResumeParsingError("No extractable text found (scanned/image-only resume?).")
    return text


def _mark_failed(application_id: int, message: str) -> None:
    Application.objects.filter(pk=application_id).update(
        parsing_status=Application.ParsingStatus.FAILED, error_message=message[:2000]
    )


@shared_task(
    bind=True,
    acks_late=True,
    autoretry_for=(OperationalError,),  # transient DB issues only
    retry_backoff=True,
    retry_jitter=True,
    retry_kwargs={"max_retries": 3},
)
def process_resume(self, application_id: int) -> dict:
    """Parse the resume, extract skills, score it against the job, persist results.

    Idempotent: completed applications are skipped, so redelivery is harmless.
    Parsing/NLP failures mark the application FAILED instead of retrying forever.
    """
    try:
        with transaction.atomic():
            application = (
                Application.objects.select_for_update()
                .select_related("job")
                .get(pk=application_id)
            )
            if application.parsing_status == Application.ParsingStatus.COMPLETED:
                return {"application_id": application_id, "skipped": True}
            application.parsing_status = Application.ParsingStatus.PROCESSING
            application.error_message = ""
            application.save(update_fields=["parsing_status", "error_message", "updated_at"])
    except Application.DoesNotExist:
        logger.warning("process_resume: application %s no longer exists", application_id)
        return {"application_id": application_id, "missing": True}

    try:
        resume_text = extract_text_from_resume(application.resume)
        score = calculate_match_score(resume_text, application.job.full_text)
        skills = extract_skills(resume_text)
    except OperationalError:
        raise  # let autoretry handle it
    except ResumeParsingError as exc:
        logger.info("Resume parsing failed for %s: %s", application_id, exc)
        _mark_failed(application_id, str(exc))
        return {"application_id": application_id, "status": "failed"}
    except Exception as exc:  # includes SoftTimeLimitExceeded
        logger.exception("Unexpected error processing application %s", application_id)
        _mark_failed(application_id, f"Processing error: {exc}")
        return {"application_id": application_id, "status": "failed"}

    Application.objects.filter(pk=application_id).update(
        match_score=score,
        extracted_skills=skills,
        parsing_status=Application.ParsingStatus.COMPLETED,
        error_message="",
    )
    return {"application_id": application_id, "status": "completed", "match_score": score}
